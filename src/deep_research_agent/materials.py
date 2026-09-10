"""Local bytes to citable text, with explicit extraction coverage."""

from __future__ import annotations

import asyncio
import csv
import io
import json
import os
import subprocess
import sys
from datetime import date, datetime, time
from pathlib import Path

from .domain import encode

TEXT_SUFFIXES = {".txt", ".md", ".json", ".yaml", ".yml", ".html"}
SUPPORTED_SUFFIXES = TEXT_SUFFIXES | {".csv", ".tsv", ".pdf", ".xlsx"}


async def parse_isolated(name: str, raw: bytes, timeout: float = 60) -> dict:
    """Run native parsers outside the host; cancellation always reaps the child."""
    if timeout <= 0:
        raise ValueError("positive parse timeout required")
    # Only OS/runtime bootstrap settings; provider credentials are not inherited.
    allowed = {"SYSTEMROOT", "WINDIR", "PATH", "TEMP", "TMP", "LANG", "LC_ALL"}
    environment = {k: v for k, v in os.environ.items() if k.upper() in allowed}
    process = await asyncio.create_subprocess_exec(
        sys.executable,
        "-m",
        "deep_research_agent.materials",
        Path(name).name,
        stdin=asyncio.subprocess.PIPE,
        stdout=asyncio.subprocess.PIPE,
        stderr=asyncio.subprocess.DEVNULL,
        env=environment,
        creationflags=subprocess.CREATE_NO_WINDOW if os.name == "nt" else 0,
    )
    try:
        try:
            output, _ = await asyncio.wait_for(process.communicate(raw), timeout)
        except TimeoutError:
            raise ValueError("material parsing timed out") from None
        if process.returncode:
            raise ValueError("material parser process failed")
        try:
            response = json.loads(output)
        except (ValueError, UnicodeError):
            raise ValueError("material parser returned invalid output") from None
        if "error" in response:
            raise ValueError(response["error"])
        return response["result"]
    finally:
        if process.returncode is None:
            try:
                process.kill()
            except ProcessLookupError:
                pass
        await process.wait()


def parse(name: str, raw: bytes) -> dict:
    suffix = Path(name).suffix.lower()
    if suffix not in SUPPORTED_SUFFIXES:
        raise ValueError("unsupported material format")
    parts, segments, issues = [], [], []
    position = 0

    def append(text, locator, status="extracted"):
        nonlocal position
        block = text + "\n"
        segments.append(
            {
                "start": position,
                "end": position + len(block),
                "locator": locator,
                "status": status,
            }
        )
        position += len(block)
        parts.append(block)

    if suffix in TEXT_SUFFIXES:
        text = raw.decode("utf-8-sig")
        return {
            "text": text,
            "segments": [
                {
                    "start": 0,
                    "end": len(text),
                    "locator": {"document": name},
                    "status": "extracted",
                }
            ],
            "coverage": "text_extracted_not_reviewed",
            "issues": [],
            "parser": "utf8-v1",
        }
    if suffix in {".csv", ".tsv"}:
        reader = csv.reader(
            io.StringIO(raw.decode("utf-8-sig"), newline=""),
            delimiter="," if suffix == ".csv" else "\t",
            strict=True,
        )
        try:
            for row, cells in enumerate(reader, 1):
                append(encode(cells), {"row": row})
        except csv.Error:
            raise ValueError("invalid delimited table") from None
        parser = "csv-v1"
    elif suffix == ".pdf":
        from pypdf import PdfReader, __version__

        try:
            reader = PdfReader(io.BytesIO(raw))
            if reader.is_encrypted:
                raise ValueError(
                    "encrypted PDF requires an unlocked user-provided copy"
                )
            for number, page in enumerate(reader.pages, 1):
                text = page.extract_text() or ""
                status = (
                    "text_extracted" if text.strip() else "needs_ocr_or_visual_review"
                )
                append(text, {"page": number}, status)
                if not text.strip():
                    issues.append({"page": number, "reason": status})
        except ValueError:
            raise
        except Exception as exc:
            raise ValueError("PDF extraction failed: " + type(exc).__name__) from None
        parser = "pypdf-" + __version__
        issues.append(
            {"reason": "text_layer_only; layout_images_and_tables_not_verified"}
        )
    else:
        from openpyxl import __version__, load_workbook

        book = None
        try:
            book = load_workbook(
                io.BytesIO(raw), read_only=True, data_only=False, keep_links=False
            )
            for sheet in book.worksheets:
                sheet.reset_dimensions()
                for row_number, row in enumerate(sheet.iter_rows(), 1):
                    cells = []
                    for cell in row:
                        if cell.value is None:
                            continue
                        value = cell.value
                        if isinstance(value, (datetime, date, time)):
                            value = value.isoformat()
                        if not isinstance(value, (str, int, float, bool)):
                            value = str(value)
                        cells.append(
                            {
                                "cell": cell.coordinate,
                                "value": value,
                                "type": cell.data_type,
                                "number_format": cell.number_format,
                            }
                        )
                    if cells:
                        append(
                            encode(cells),
                            {
                                "sheet": sheet.title,
                                "row": row_number,
                                "visibility": sheet.sheet_state,
                            },
                        )
            issues.append(
                {
                    "reason": "formulas_not_evaluated; charts_images_and_merged_layout_not_extracted"
                }
            )
        except Exception as exc:
            raise ValueError("XLSX extraction failed: " + type(exc).__name__) from None
        finally:
            if book is not None:
                book.close()
        parser = "openpyxl-" + __version__
    return {
        "text": "".join(parts),
        "segments": segments,
        "coverage": "partial_extraction" if issues else "table_extracted_not_reviewed",
        "issues": issues,
        "parser": parser,
    }


if __name__ == "__main__":
    try:
        result = {"result": parse(sys.argv[1], sys.stdin.buffer.read())}
    except Exception as exc:
        result = {"error": "material extraction failed: " + type(exc).__name__}
    sys.stdout.buffer.write(json.dumps(result, ensure_ascii=False).encode("utf-8"))
