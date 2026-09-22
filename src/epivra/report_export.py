"""Deterministic Word export of the Host's published Markdown, without network I/O."""

from io import BytesIO

from docx import Document
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor
from markdown_it import MarkdownIt

from .markdown_rules import math_plugin


def markdown_report(report):
    # Only generated references are literal; author numeric links retain their meaning.
    text, end = report["text"], 0
    parts: list[str] = []
    for mark in report.get("citation_marks", []):
        parts.extend((text[end:mark["start"]], "\\" + text[mark["start"]:mark["end"]]))
        end = mark["end"]
    parts.append(text[end:])
    return "".join(parts)


def word_report(report):
    document = Document()
    document.core_properties.identifier = report["ref"]
    document.core_properties.title = "Epivra research report"
    section = document.sections[0]
    section.page_width, section.page_height = Cm(21), Cm(29.7)
    section.top_margin = section.bottom_margin = Cm(2)
    section.left_margin = section.right_margin = Cm(2)
    normal = document.styles["Normal"]
    normal.font.name = "Calibri"
    normal.font.size = Pt(11)
    normal.element.get_or_add_rPr().rFonts.set(qn("w:eastAsia"), "SimSun")
    normal.paragraph_format.space_after = Pt(8)
    normal.paragraph_format.line_spacing = 1.25
    for level in range(1, 7):
        style = document.styles[f"Heading {level}"]
        style.font.color.rgb = RGBColor.from_string("1D5545")
        style.element.get_or_add_rPr().rFonts.set(qn("w:eastAsia"), "Microsoft YaHei")
    parser = math_plugin(MarkdownIt("default", {"html": False}))
    tokens = parser.parse(markdown_report(report))
    paragraph = None
    table = None
    row = -1
    cell = -1
    lists = []
    quote_depth = 0

    def add_inline(target, children):
        bold = italic = strike = False
        href = None
        for token in children:
            if token.type in {"strong_open", "strong_close"}:
                bold = token.type.endswith("open")
            elif token.type in {"em_open", "em_close"}:
                italic = token.type.endswith("open")
            elif token.type in {"s_open", "s_close"}:
                strike = token.type.endswith("open")
            elif token.type == "link_open":
                href = token.attrGet("href")
            elif token.type == "link_close":
                if href:
                    target.add_run(f" ({href})")
                href = None
            elif token.type in {"softbreak", "hardbreak"}:
                target.add_run().add_break()
            elif token.type == "image":
                target.add_run(f"[{token.content}] ({token.attrGet('src')})")
            elif token.type in {"text", "code_inline", "html_inline", "math_inline"}:
                run = target.add_run(token.content)
                run.bold, run.italic, run.font.strike = bold, italic, strike
                if token.type in {"code_inline", "math_inline"}:
                    run.font.name = "Consolas"
                if href:
                    run.font.color.rgb = RGBColor.from_string("1D5545")

    for index, token in enumerate(tokens):
        kind = token.type
        if kind == "heading_open":
            paragraph = document.add_heading(level=int(token.tag[1:]))
        elif kind in {"bullet_list_open", "ordered_list_open"}:
            lists.append(
                {
                    "number": int(token.attrGet("start") or 1)
                    if kind == "ordered_list_open"
                    else None,
                    "first": False,
                }
            )
        elif kind in {"bullet_list_close", "ordered_list_close"}:
            lists.pop()
        elif kind == "list_item_open":
            lists[-1]["first"] = True
        elif kind == "blockquote_open":
            quote_depth += 1
        elif kind == "blockquote_close":
            quote_depth -= 1
        elif kind == "table_open":
            count = 0
            for following in tokens[index + 1 :]:
                if following.type == "tr_close":
                    break
                if following.type == "th_open":
                    count += 1
            table = document.add_table(rows=0, cols=count)
            table.style = "Table Grid"
            row = -1
        elif kind == "table_close":
            table = None
        elif kind == "tr_open":
            table.add_row()
            row += 1
            cell = -1
            if row == 0:
                table.rows[0]._tr.get_or_add_trPr().append(OxmlElement("w:tblHeader"))
        elif kind in {"td_open", "th_open"}:
            cell += 1
            paragraph = table.cell(row, cell).paragraphs[0]
            alignment = token.attrGet("style") or ""
            if "right" in alignment:
                paragraph.alignment = 2
            elif "center" in alignment:
                paragraph.alignment = 1
        elif kind == "paragraph_open":
            paragraph = document.add_paragraph()
            if lists:
                if lists[-1]["first"]:
                    number = lists[-1]["number"]
                    paragraph.add_run(f"{number}. " if number is not None else "• ")
                    if number is not None:
                        lists[-1]["number"] += 1
                    lists[-1]["first"] = False
            if quote_depth or lists:
                paragraph.paragraph_format.left_indent = Cm(0.6 * quote_depth + 0.5 * len(lists))
        elif kind == "inline":
            add_inline(paragraph, token.children or [])
            if table is not None and row == 0:
                for run in paragraph.runs:
                    run.bold = True
        elif kind in {"fence", "code_block", "math_block"}:
            p = document.add_paragraph()
            run = p.add_run(token.content.rstrip("\n"))
            run.font.name, run.font.size = "Consolas", Pt(9)
        elif kind == "hr":
            document.add_paragraph("—" * 20)
    output = BytesIO()
    document.save(output)
    return output.getvalue()
