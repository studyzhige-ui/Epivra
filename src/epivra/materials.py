"""Local bytes to citable text, with explicit extraction coverage."""

from __future__ import annotations

import asyncio
import csv
import io
import json
import os
import subprocess
import sys
import weakref
import zipfile
from copy import deepcopy
from datetime import date, datetime, time
from pathlib import Path

from .document_parser import FORMATS
from .domain import encode

TEXT_SUFFIXES = {".txt", ".md", ".json", ".yaml", ".yml", ".html"}
SUPPORTED_SUFFIXES = TEXT_SUFFIXES | {".csv", ".tsv", ".pdf", ".xlsx"} | FORMATS
MAX_INPUT_BYTES = 256 * 1024 * 1024
MAX_OUTPUT_BYTES = 64 * 1024 * 1024
MAX_ELEMENTS = 1_000_000
_parse_slots = weakref.WeakKeyDictionary()


def _bounded(result):
    if len(encode({"result": result}).encode("utf-8")) > MAX_OUTPUT_BYTES:
        raise ValueError("material extraction exceeds parser capacity")
    return result


async def parse_isolated(
    name: str, raw: bytes, timeout: float = 60, options=None
) -> dict:
    if len(raw) > MAX_INPUT_BYTES:
        raise ValueError("material exceeds input byte limit")
    loop = asyncio.get_running_loop()
    gate = _parse_slots.setdefault(loop, asyncio.Semaphore(2))
    async with gate:
        return await _parse_process(name, raw, timeout, options)


async def _parse_process(name, raw, timeout, options):
    """Run native parsers outside the host; cancellation always reaps the child."""
    if timeout <= 0:
        raise ValueError("positive parse timeout required")
    # Only OS/runtime bootstrap settings; provider credentials are not inherited.
    allowed = {
        "SYSTEMROOT",
        "WINDIR",
        "PATH",
        "TEMP",
        "TMP",
        "LANG",
        "LC_ALL",
        "HOME",
        "USERPROFILE",
        "HOMEDRIVE",
        "HOMEPATH",
    }
    environment = {k: v for k, v in os.environ.items() if k.upper() in allowed}
    environment.update(
        HF_HUB_OFFLINE="1", TRANSFORMERS_OFFLINE="1", HF_HUB_DISABLE_TELEMETRY="1"
    )
    process = await asyncio.create_subprocess_exec(
        sys.executable,
        "-m",
        "epivra.materials",
        Path(name).name,
        encode(options or {}),
        stdin=asyncio.subprocess.PIPE,
        stdout=asyncio.subprocess.PIPE,
        stderr=asyncio.subprocess.DEVNULL,
        env=environment,
        creationflags=subprocess.CREATE_NO_WINDOW if os.name == "nt" else 0,
    )

    async def feed():
        try:
            for offset in range(0, len(raw), 65536):
                process.stdin.write(raw[offset : offset + 65536])
                await process.stdin.drain()
        except (BrokenPipeError, ConnectionResetError):
            pass
        finally:
            process.stdin.close()

    async def read():
        output = bytearray()
        while chunk := await process.stdout.read(65536):
            if len(output) + len(chunk) > MAX_OUTPUT_BYTES:
                raise ValueError("material parser exceeds output byte limit")
            output.extend(chunk)
        return output

    jobs = [asyncio.create_task(feed()), asyncio.create_task(read())]
    try:
        try:
            async with asyncio.timeout(timeout):
                _, output = await asyncio.gather(*jobs)
                await process.wait()
        except TimeoutError:
            raise ValueError("material parsing timed out") from None
        if process.returncode:
            raise ValueError("material parser process failed")
        try:
            response = json.loads(output)
        except (ValueError, UnicodeError):
            raise ValueError("material parser returned invalid output") from None
        if "error" in response:
            raise ValueError(response["error"])
        return response["result"]
    finally:
        for job in jobs:
            job.cancel()
        if process.returncode is None:
            try:
                process.kill()
            except ProcessLookupError:
                pass

        async def reap():
            await asyncio.gather(*jobs, return_exceptions=True)
            # A killed child can still have a paused, full stdout pipe. Drain it
            # without retaining bytes before waiting for transport completion.
            while await process.stdout.read(65536):
                pass
            await process.wait()

        cleanup = asyncio.create_task(reap())
        cancelled = False
        while not cleanup.done():
            try:
                await asyncio.shield(cleanup)
            except asyncio.CancelledError:
                # Repeated cancellation must not leave the parser unreaped.
                cancelled = True
        cleanup.result()
        if cancelled:
            raise asyncio.CancelledError


def parse(name: str, raw: bytes, options=None) -> dict:
    if len(raw) > MAX_INPUT_BYTES:
        raise ValueError("material exceeds input byte limit")
    if zipfile.is_zipfile(io.BytesIO(raw)):
        with zipfile.ZipFile(io.BytesIO(raw)) as archive:
            entries = archive.infolist()
            if (
                len(entries) > MAX_ELEMENTS
                or sum(x.file_size for x in entries) > MAX_INPUT_BYTES
            ):
                raise ValueError("document expanded content exceeds parser capacity")
    options = options or {}
    mode = options.get("parser", "auto")
    suffix = Path(name).suffix.lower()
    if suffix not in SUPPORTED_SUFFIXES:
        raise ValueError("unsupported material format")
    if mode == "light" and suffix in FORMATS:
        raise ValueError("this format needs Docling; select auto or docling parsing")
    if (mode == "docling" and suffix == ".pdf") or suffix in FORMATS:
        from .document_parser import convert

        return _bounded(convert(name, raw, options))
    parts, segments, issues = [], [], []
    position = 0
    output_bytes = 0
    encoding = options.get("encoding", "utf-8-sig")

    def append(text, locator, status="extracted"):
        nonlocal position, output_bytes
        block = text + "\n"
        segment = {"start": position, "end": position + len(block), "locator": locator, "status": status}
        output_bytes += len(encode(block)[1:-1].encode("utf-8")) + len(encode(segment).encode("utf-8")) + 1
        if len(segments) >= MAX_ELEMENTS or output_bytes > MAX_OUTPUT_BYTES:
            raise ValueError("material extraction exceeds parser capacity")
        segments.append(
            {
                "start": position,
                "end": position + len(block),
                "locator": locator,
                "status": status,
            }
        )
        position += len(block)
        parts.append(block)

    if suffix in TEXT_SUFFIXES:
        text = raw.decode(encoding)
        result = {
            "text": text,
            "segments": [
                {
                    "start": 0,
                    "end": len(text),
                    "locator": {"document": name},
                    "status": "extracted",
                }
            ],
            "coverage": "text_extracted_not_reviewed",
            "issues": [],
            "parser": "text-" + encoding,
        }
        return _bounded(result)
    if suffix in {".csv", ".tsv"}:
        reader = csv.reader(
            io.StringIO(raw.decode(encoding), newline=""),
            delimiter="," if suffix == ".csv" else "\t",
            strict=True,
        )
        try:
            for row, cells in enumerate(reader, 1):
                append(encode(cells), {"row": row})
        except csv.Error:
            raise ValueError("invalid delimited table") from None
        parser = "csv-v1"
    elif suffix == ".pdf":
        from pypdf import PdfReader, __version__

        try:
            reader = PdfReader(io.BytesIO(raw))
            if reader.is_encrypted:
                raise ValueError(
                    "encrypted PDF requires an unlocked user-provided copy"
                )
            for number, page in enumerate(reader.pages, 1):
                text = page.extract_text() or ""
                status = (
                    "text_extracted" if text.strip() else "blank_page"
                    if page.get_contents() is None and not page.get("/Resources", {}).get("/XObject")
                    else "needs_ocr_or_visual_review"
                )
                append(text, {"page": number}, status)
                if not text.strip():
                    issues.append({"page": number, "reason": status})
        except ValueError:
            raise
        except Exception as exc:
            raise ValueError("PDF extraction failed: " + type(exc).__name__) from None
        parser = "pypdf-" + __version__
        if (
            mode == "auto"
            and options.get("artifacts_path")
            and any(i.get("reason") == "needs_ocr_or_visual_review" for i in issues)
        ):
            from .document_parser import convert

            original = parts[:], segments[:], issues[:], position, output_bytes
            try:
                from pypdf import PdfWriter

                selected = [i["page"] for i in issues if i["reason"] == "needs_ocr_or_visual_review"]
                subset = PdfWriter()
                for number in selected:
                    subset.add_page(reader.pages[number - 1])
                stream = io.BytesIO()
                subset.write(stream)
                converted = convert(name, stream.getvalue(), options)
                by_page = {}
                for segment in converted["segments"]:
                    locator = deepcopy(segment["locator"])
                    page_no = locator.get("page")
                    if page_no is None or not 1 <= page_no <= len(selected):
                        continue
                    locator["page"] = selected[page_no - 1]
                    for location in locator.get("locations", []):
                        location["page"] = selected[location["page"] - 1]
                    by_page.setdefault(locator["page"], []).append((converted["text"][segment["start"]:segment["end"]], locator))
                old_parts, old_segments = parts[:], segments[:]
                parts.clear()
                segments.clear()
                position = output_bytes = 0
                for block, segment in zip(old_parts, old_segments):
                    number = segment["locator"]["page"]
                    if number in by_page:
                        for value, locator in by_page[number]:
                            append(value, locator, "extracted_not_reviewed")
                    else:
                        append(block[:-1], segment["locator"], segment["status"])
                issues = [i for i in issues if i.get("page") not in by_page]
                for issue in deepcopy(converted.get("issues", [])):
                    locator = issue.get("locator", {})
                    if "page" in locator:
                        locator["page"] = selected[locator["page"] - 1]
                    for location in locator.get("locations", []):
                        location["page"] = selected[location["page"] - 1]
                    if "page" in issue and 1 <= issue["page"] <= len(selected):
                        issue["page"] = selected[issue["page"] - 1]
                    issues.append(issue)
                parser += "+" + converted["parser"]
            except (ValueError, RuntimeError, ImportError):
                parts, segments, issues, position, output_bytes = original
                issues.append(
                    {"reason": "docling_unavailable; retained_partial_text_layer"}
                )
        issues.append(
            {"reason": "text_layer_only; layout_images_and_tables_not_verified"}
        )
    else:
        from openpyxl import __version__, load_workbook

        book = None
        try:
            book = load_workbook(
                io.BytesIO(raw), read_only=True, data_only=False, keep_links=False
            )
            for sheet in book.worksheets:
                sheet.reset_dimensions()
                for row_number, row in enumerate(sheet.iter_rows(), 1):
                    cells = []
                    for cell in row:
                        if cell.value is None:
                            continue
                        value = cell.value
                        if isinstance(value, (datetime, date, time)):
                            value = value.isoformat()
                        if not isinstance(value, (str, int, float, bool)):
                            value = str(value)
                        cells.append(
                            {
                                "cell": cell.coordinate,
                                "value": value,
                                "type": cell.data_type,
                                "number_format": cell.number_format,
                            }
                        )
                    if cells:
                        append(
                            encode(cells),
                            {
                                "sheet": sheet.title,
                                "row": row_number,
                                "visibility": sheet.sheet_state,
                            },
                        )
            issues.append(
                {
                    "reason": "formulas_not_evaluated; charts_images_and_merged_layout_not_extracted"
                }
            )
        except Exception as exc:
            raise ValueError("XLSX extraction failed: " + type(exc).__name__) from None
        finally:
            if book is not None:
                book.close()
        parser = "openpyxl-" + __version__
    result = {
        "text": "".join(parts),
        "segments": segments,
        "coverage": "partial_extraction" if issues else "table_extracted_not_reviewed",
        "issues": issues,
        "parser": parser,
    }

    return _bounded(result)


if __name__ == "__main__":
    import socket

    def deny_network(*args, **kwargs):
        raise OSError("material parser network access is disabled")

    socket.socket.connect = deny_network
    socket.create_connection = deny_network
    try:
        result = {
            "result": parse(
                sys.argv[1],
                sys.stdin.buffer.read(),
                json.loads(sys.argv[2]) if len(sys.argv) > 2 else {},
            )
        }
    except ValueError as exc:
        result = {"error": str(exc)}
    except Exception as exc:
        result = {"error": "material extraction failed: " + type(exc).__name__}
    sys.stdout.buffer.write(encode(result).encode("utf-8"))
