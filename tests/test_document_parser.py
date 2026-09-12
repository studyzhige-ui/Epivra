import io
import tempfile
import unittest
from importlib.util import find_spec
from pathlib import Path
from unittest.mock import patch

from deep_research_agent.materials import parse, parse_isolated
from deep_research_agent.storage import Store
from deep_research_agent.workspace import Workspace


class DocumentParserTests(unittest.IsolatedAsyncioTestCase):
    @unittest.skipUnless(
        find_spec("docling") and find_spec("docx"), "optional documents extra"
    )
    async def test_actual_isolated_docx_chinese_table_and_located_text(self):
        from docx import Document

        doc = Document()
        doc.add_heading("原始研究资料", 0)
        doc.add_paragraph("只有满足条件时才适用。")
        table = doc.add_table(rows=2, cols=2)
        for cell, value in zip(
            [c for r in table.rows for c in r.cells],
            ["年份", "数量", "2025", "42"],
            strict=True,
        ):
            cell.text = value
        output = io.BytesIO()
        doc.save(output)
        body = await parse_isolated(
            "资料.docx", output.getvalue(), 90, {"parser": "auto"}
        )
        self.assertIn("只有满足条件时才适用。", body["text"])
        self.assertIn("2025", body["text"])
        self.assertIn("42", body["text"])
        self.assertTrue(body["parser"].startswith("docling-"))
        self.assertTrue(all(s["locator"]["element"] for s in body["segments"]))
        self.assertTrue(
            all(body["text"][s["start"] : s["end"]].strip() for s in body["segments"])
        )

    async def test_identical_bytes_under_new_name_reuse_extraction(self):
        with tempfile.TemporaryDirectory() as folder:
            store = Store(Path(folder) / "state.db")
            try:
                control = store.create("s", "Research", {})
                workspace = Workspace(store)
                await workspace.upload_async("s", control.ref, "first.txt", b"original")
                with patch(
                    "deep_research_agent.workspace.parse_isolated",
                    side_effect=AssertionError("must reuse"),
                ):
                    second = await workspace.upload_async(
                        "s", control.ref, "renamed.txt", b"original"
                    )
                self.assertEqual("original", second.body["text"])
                self.assertEqual(1, len(store.list("s", "material_bytes")))
            finally:
                store.close()

    async def test_explicit_light_and_missing_ocr_resources_fail_clearly(self):
        with self.assertRaisesRegex(ValueError, "needs Docling"):
            parse("test.docx", b"unused", {"parser": "light"})
        if find_spec("docling"):
            with self.assertRaisesRegex(ValueError, "model directory"):
                await parse_isolated("scan.pdf", b"unused", 90, {"parser": "docling"})
