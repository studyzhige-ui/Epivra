import tempfile
import unittest
from pathlib import Path

from epivra.citations import render, validate
from epivra.storage import Store


class CitationTests(unittest.TestCase):
    def setUp(self):
        self.temp = tempfile.TemporaryDirectory()
        self.store = Store(Path(self.temp.name) / "store.db")
        self.store.create("s", "Question", {})
        self.a = self.store.put(
            "s", "source", {"text": "First passage", "origin": "https://example.org/a"}
        )
        self.b = self.store.put(
            "s", "source", {"text": "Second passage", "origin": "notes.txt"}
        )
        self.resolve = lambda ref: self.store.get("s", ref)

    def tearDown(self):
        self.store.close()
        self.temp.cleanup()

    def report(self, text, evidence=None):
        evidence = evidence if evidence is not None else [self.a.ref, self.b.ref]
        result = {**render(text, evidence, self.resolve), "evidence": evidence}
        validate(result, self.resolve)
        return result

    def test_historical_math_and_link_citations_remain_readable(self):
        from epivra.presentation import published_report
        for text in (f"Formula $x[[cite:{self.a.ref}]]$.", f"[label [[cite:{self.a.ref}]]](https://example.com)"):
            old = render(text, [self.a.ref], self.resolve, _historical=True)
            old.pop("citation_marks")
            old["evidence"] = [self.a.ref]
            validate(old, self.resolve)
            report = self.store.put("s", "report", old)
            direction = self.store.control("s").direction
            self.store._put("s", "publication", {"report": report.ref}, (direction, report.ref))
            view = published_report(self.store, "s")
            self.assertEqual([], view["citation_marks"])
            self.assertEqual(old["text"], view["text"])
            with self.assertRaises(ValueError):
                render(text, [self.a.ref], self.resolve)

    def test_first_appearance_not_evidence_order_and_repeat(self):
        report = self.report(
            f"B [[cite:{self.b.ref}]]. A [[cite:{self.a.ref}]]. B [[cite:{self.b.ref}]]."
        )
        self.assertTrue(report["text"].startswith("B [1]. A [2]. B [1]."))
        self.assertIn("1. notes.txt", report["text"])
        self.assertNotIn("#source-", report["text"])
        self.assertEqual([self.b.ref, self.a.ref, self.b.ref], report["citations"])

    def test_saved_title_is_used_without_changing_source_binding(self):
        source = self.store.put(
            "s",
            "source",
            {
                "text": "Evidence",
                "title": "Research [2026]",
                "origin": "https://example.org/paper",
            },
        )
        report = self.report(f"Finding [[cite:{source.ref}]].", [source.ref])
        self.assertIn(
            r"[Research \[2026\]](<https://example.org/paper>)", report["text"]
        )
        self.assertEqual([source.ref], report["citations"])

    def test_source_and_two_passages_share_source_number(self):
        note = self.store.put(
            "s",
            "note",
            {"source": self.a.ref, "quote": "First", "offset": 0},
            (self.a.ref,),
        )
        note2 = self.store.put(
            "s",
            "note",
            {"source": self.a.ref, "quote": "passage", "offset": 6},
            (self.a.ref,),
        )
        report = self.report(
            f"One [[cite:{note.ref}]] two [[cite:{note2.ref}]] all [[cite:{self.a.ref}]]"
        )
        self.assertTrue(report["text"].startswith("One [1] two [1] all [1]"))

    def test_zero_citations_and_markdown_code_are_preserved(self):
        for text in [
            "Pure reasoning.",
            "    arr = [1]\n",
            "~~~python\narr = [1]\n~~~~\n",
            "```python\narr = [1]",
            "`[1]` and ``a `[2]` b``",
            "\\[1]",
            "Code:\n\n> ```\n> [1]\n> ```",
        ]:
            with self.subTest(text=text):
                self.assertEqual(text, self.report(text, [])["text"])

    def test_unknown_and_malformed_and_manual_references_are_rejected(self):
        for text in [
            "[[cite:" + "a" * 64 + "]]",
            "[[cite:no]]",
            "Lone ` opener.\n\n[[cite:" + "a" * 64 + "]].\n\nLone ` closer.",
        ]:
            with self.subTest(text=text), self.assertRaises(ValueError):
                self.report(text)

    def test_wrong_evidence_and_bad_passage(self):
        with self.assertRaises(ValueError):
            self.report(f"[[cite:{self.a.ref}]]", [self.b.ref])
        note = self.store.put(
            "s",
            "note",
            {"source": self.a.ref, "quote": "First", "offset": 1},
            (self.a.ref,),
        )
        with self.assertRaises(ValueError):
            self.report(f"[[cite:{note.ref}]]")
        other = self.store.create("other", "Other", {})
        source = self.store.put(
            "other", "source", {"text": "Foreign"}, (other.direction,)
        )
        with self.assertRaises(ValueError):
            self.report(f"[[cite:{source.ref}]]")

    def test_bibliography_cannot_be_swallowed_by_unclosed_code(self):
        with self.assertRaisesRegex(ValueError, "Close the fenced"):
            self.report(f"Claim [[cite:{self.a.ref}]]\n\n```python\n[1]")
        report = self.report(f"Claim [[cite:{self.a.ref}]]\n\n~~~python\n[1]\n~~~~")
        self.assertTrue(report["text"].startswith("Claim [1]"))

    def test_link_backtick_does_not_mask_visible_invalid_citation(self):
        text = "[Docs](https://example.org/`) Claim [[cite:bad]] `sample`"
        with self.assertRaises(ValueError):
            self.report(text, [])
        with self.assertRaises(ValueError):
            validate({"text": text, "evidence": []}, self.resolve)

    def test_container_can_implicitly_end_fence_before_bibliography(self):
        for prefix in ("> ```\n> sample code\n\n", "- ```\n  sample code\n\n"):
            with self.subTest(prefix=prefix):
                report = self.report(prefix + f"Claim [[cite:{self.a.ref}]]")
                self.assertTrue(report["text"].startswith(prefix + "Claim [1]"))

    def test_preserves_locations_with_containers_code_links_and_repeated_markers(self):
        marker = f"[[cite:{self.a.ref}]]"
        for text in (
            f"> `{marker}` then {marker}",
            f"- [Docs](https://example.org/`) {marker} `sample`",
            f"## Heading {marker}\r\n\r\n> Quote {marker}",
        ):
            with self.subTest(text=text):
                report = self.report(text)
                self.assertIn("[1]", report["text"])

    def test_commonmark_newline_mapping_preserves_original_characters(self):
        for separator in ("\r", "\r\n", "\n", "\u2028", "\v", "\u0085"):
            with self.subTest(separator=repr(separator)):
                report = self.report(f"Text{separator}Claim [[cite:{self.a.ref}]]")
                self.assertTrue(report["text"].startswith(f"Text{separator}Claim [1]"))

    def test_html_is_literal_like_the_web_report_renderer(self):
        with self.assertRaises(ValueError):
            self.report("<div>Claim [[cite:bad]]</div>", [])

    def test_table_cells_use_parser_scopes_and_original_positions(self):
        marker = f"[[cite:{self.a.ref}]]"
        text = f"| Code | Finding |\n|---|---|\n| `{marker}` | {marker} |"
        report = self.report(text)
        self.assertTrue(
            report["text"].startswith(text.replace(f"| {marker} |", "| [1] |"))
        )
        with self.assertRaises(ValueError):
            self.report("| A | B |\n|---|---|\n| `literal | [[cite:bad]] ` |", [])

    def test_numeric_links_coexist_with_bound_citations_in_word(self):
        from io import BytesIO

        from docx import Document

        from epivra.report_export import word_report

        text = f"See [1]. Claim [[cite:{self.a.ref}]].\n\n[1]: https://author.example"
        report = self.report(text)
        self.assertEqual(1, len(report["citation_marks"]))
        output = Document(BytesIO(word_report({"ref": "r", **report})))
        paragraph = output.paragraphs[0].text
        self.assertEqual("See 1 (https://author.example). Claim [1].", paragraph)
        self.assertEqual("See [1].\n\n[1]: https://author.example",
                         self.report("See [1].\n\n[1]: https://author.example", [])["text"])

    def test_math_container_contracts_preserve_outside_citations(self):
        import json
        from io import BytesIO

        from docx import Document
        from markdown_it import MarkdownIt

        from epivra.markdown_rules import math_plugin
        from epivra.report_export import word_report

        cases = json.loads((Path(__file__).parent / "fixtures/markdown_contracts.json").read_text(encoding="utf-8"))
        for case in cases:
            with self.subTest(text=case["text"]):
                text = case["text"].replace("@CITE@", f"[[cite:{self.a.ref}]]")
                report = self.report(text)
                self.assertEqual(1, len(report["citation_marks"]))
                tokens = math_plugin(MarkdownIt()).parse(text)
                self.assertEqual(case["math_blocks"], sum(t.type == "math_block" for t in tokens))
                self.assertFalse(any("Outside" in t.content for t in tokens if t.type == "math_block"))
                output = Document(BytesIO(word_report({"ref": "r", **report})))
                outside = [p for p in output.paragraphs if "Outside" in p.text]
                self.assertEqual(1, len(outside))
                self.assertEqual("Outside [1]", outside[0].text)
                self.assertTrue(all(r.font.name != "Consolas" for r in outside[0].runs))

    def test_publication_revalidation_rejects_tampered_rendering(self):
        report = self.report(f"Claim [[cite:{self.a.ref}]]")
        for change in [
            {"text": report["text"].replace("[1]", "[2]", 1)},
            {"citations": [self.b.ref]},
            {"text": report["text"] + "\nTampered bibliography"},
            {"citations": []},
            {"evidence": []},
        ]:
            with self.subTest(change=change), self.assertRaises(ValueError):
                validate({**report, **change}, self.resolve)

    def test_ordinary_brackets_are_not_citation_bindings(self):
        text = f"Year [2024], range [0, 1], weights[0], claim [[cite:{self.a.ref}]] then [1]."
        report = self.report(text)
        self.assertEqual(1, len(report["citation_marks"]))
        self.assertTrue(report["text"].startswith("Year [2024], range [0, 1], weights[0], claim [1] then [1]."))

    def test_citation_placement_respects_math_and_links(self):
        marker = f"[[cite:{self.a.ref}]]"
        for text in (f"$x{marker}$", f"[Label {marker}](https://example.org)", f"$$\nx{marker}\n$$"):
            with self.subTest(text=text), self.assertRaisesRegex(ValueError, "outside"):
                self.report(text)
