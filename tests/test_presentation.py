"""Public projections and deterministic exports; no research API calls."""

import tempfile
import unittest
from io import BytesIO
from pathlib import Path

from docx import Document

from epivra.citations import render
from epivra.domain import Conflict
from epivra.presentation import progress, published_report, source_page, work_detail
from epivra.report_export import word_report
from epivra.storage import Store


class PresentationTests(unittest.TestCase):
    def setUp(self):
        self.folder = tempfile.TemporaryDirectory()
        self.store = Store(Path(self.folder.name) / "research.db")
        self.c = self.store.create("s", "中文研究", {})

    def tearDown(self):
        self.store.close()
        self.folder.cleanup()

    def work(self, task, role="investigator"):
        return self.store._put(
            "s",
            "work",
            {"role": role, "task": task, "direction": self.c.direction},
            (self.c.direction,),
        )

    def test_citations_bind_quotes_and_source_pages_to_exact_snapshots(self):
        source = self.store.put(
            "s",
            "source",
            {
                "text": "前文。精确摘录。后文。",
                "origin": "https://example.org/a",
                "coverage": "partial_extraction",
            },
        )
        note = self.store.put(
            "s",
            "note",
            {"source": source.ref, "quote": "精确摘录", "offset": 3},
            (source.ref,),
        )
        unused = self.store.put("s", "source", {"text": "其他依据"})
        body = render(
            f"结论 [[cite:{note.ref}]]；再引 [[cite:{source.ref}]]。",
            [source.ref, unused.ref],
            lambda ref: self.store.get("s", ref),
        )
        report = self.store.put(
            "s",
            "report",
            {**body, "evidence": [source.ref, unused.ref]},
            (self.c.direction, source.ref, unused.ref),
        )
        self.store._put(
            "s", "publication", {"report": report.ref}, (self.c.direction, report.ref)
        )
        view = published_report(self.store, "s", report.ref)
        self.assertEqual(view["text"], body["text"])
        self.assertEqual(len(view["sources"]), 2)
        self.assertEqual(
            view["citations"],
            [
                {
                    "number": 1,
                    "source": source.ref,
                    "quotes": [{"text": "精确摘录", "offset": 3}],
                }
            ],
        )
        page = source_page(self.store, "s", source.ref, 3, 4)
        self.assertEqual(page["text"], "精确摘录")
        self.assertEqual(page["coverage"], "partial_extraction")
        self.assertEqual(page["next_offset"], 7)
        for offset in (-1, True):
            with self.assertRaises(ValueError):
                source_page(self.store, "s", source.ref, offset)
        with self.assertRaises(ValueError):
            source_page(self.store, "s", note.ref)
        with self.assertRaises(Conflict):
            published_report(self.store, "s", "old-report")
        self.store.command("s", "steer", self.c.ref, "steer", {"request": "new"})
        self.assertEqual(published_report(self.store, "s"), {"report": None})
        with self.assertRaises(Conflict):
            published_report(self.store, "s", report.ref)

    def test_waits_do_not_mark_children_waiting_and_delivery_is_not_verification(self):
        lead, child = self.work("协调", "lead"), self.work("调查")
        self.store.put(
            "s",
            "work_wait",
            {"producer": lead.ref, "refs": [child.ref]},
            (lead.ref, child.ref),
        )
        self.store.put("s", "step", {"reasoning_content": "PRIVATE"}, (child.ref,))
        self.store.put(
            "s", "note", {"producer": child.ref, "text": "公开发现"}, (child.ref,)
        )
        states = {w["ref"]: w["state"] for w in progress(self.store, "s")["work"]}
        self.assertEqual(states, {lead.ref: "waiting", child.ref: "pending"})
        result = self.store.put(
            "s",
            "work_result",
            {"producer": child.ref, "text": "还缺独立数据", "refs": []},
            (child.ref,),
        )
        views = {w["ref"]: w for w in progress(self.store, "s")["work"]}
        self.assertEqual(views[lead.ref]["state"], "pending")
        self.assertEqual(views[child.ref]["state"], "delivered")
        self.assertEqual(views[child.ref]["result"], result.ref)
        detail = work_detail(self.store, "s", child.ref)
        self.assertNotIn("PRIVATE", str(detail))
        self.assertEqual(
            [e["text"] for e in detail["entries"]], ["公开发现", "还缺独立数据"]
        )
        self.store.command("s", "steer", self.c.ref, "steer", {"request": "new"})
        self.assertEqual(progress(self.store, "s")["work"], [])
        with self.assertRaises(Conflict):
            work_detail(self.store, "s", child.ref)

    def test_word_preserves_nested_lists_chinese_table_code_and_source_urls(self):
        report = {
            "ref": "report:version",
            "text": "# 中文标题\n\n正文 **重点** 与 [来源](https://example.org)。\n\n3. 第三项\n   - 内层\n4. 第四项\n\n| 方案 | 数值 |\n|---|---:|\n|甲|12|\n\n```python\nprint('<script>')\n```\n\n![原图说明](https://example.org/image.png)",
        }
        doc = Document(BytesIO(word_report(report)))
        text = "\n".join(p.text for p in doc.paragraphs)
        self.assertIn("中文标题", text)
        self.assertIn("https://example.org", text)
        self.assertIn("3. 第三项", text)
        self.assertIn("4. 第四项", text)
        self.assertIn("• 内层", text)
        self.assertIn("print('<script>')", text)
        self.assertIn("原图说明", text)
        self.assertEqual(doc.tables[0].cell(1, 1).text, "12")
        self.assertEqual(doc.core_properties.identifier, report["ref"])

    def test_word_keeps_formula_delimiters_and_code_literal(self):
        text = r"Formula \(x^2\), \[y=x\], $a_b$ and `\(literal\)`"
        doc = Document(BytesIO(word_report({"ref": "v", "text": text})))
        self.assertEqual(doc.paragraphs[0].text, text.replace("`", ""))

    def test_escaped_citation_is_not_an_occurrence(self):
        source = self.store.put("s", "source", {"text": "source"})
        body = render(
            r"Literal \[1], actual " + f"[[cite:{source.ref}]]",
            [source.ref],
            lambda ref: self.store.get("s", ref),
        )
        report = self.store.put(
            "s",
            "report",
            {**body, "evidence": [source.ref]},
            (self.c.direction, source.ref),
        )
        self.store._put(
            "s", "publication", {"report": report.ref}, (self.c.direction, report.ref)
        )
        view = published_report(self.store, "s")
        self.assertEqual(len(view["citation_marks"]), 1)
        mark = view["citation_marks"][0]
        self.assertEqual(body["text"][mark["start"] : mark["end"]], "[1]")
        self.assertGreater(mark["start"], body["text"].index("actual"))
